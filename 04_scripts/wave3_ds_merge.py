# -*- coding: utf-8 -*-
"""WAVE 3 vong DATASET — GOP ket qua trich xuat + TONG HOP muc dataset + xuat XLSX.

1. Doc ds_batch_*_out.jsonl (doi chieu ds_manifest.json; chay duoc nhieu lan,
   chi can du lo nao gop lo do).
2. CAP NHAT VAO COT DA CO cua wave3_ai_final.csv (theo yeu cau Vy, khong tao cot trung):
   - dataset  <- ds_ten (ban cu luu vao dataset_cu_heuristic neu khac)
   - therapy  <- therapy_ra_soat (ban cu luu vao therapy_cu_heuristic neu khac)
   Cot MOI (chua ton tai): lien_quan_dataset, nhom_mien, phuong_thuc, don_da_luot,
   ngon_ngu, nguon_du_lieu, ds_minh_chung.
3. TONG HOP MUC DATASET (don vi phan tich): gop bai theo ten chuan hoa ->
   moi dataset 1 dong: ten, nhom_mien, phuong_thuc, don_da_luot, ngon_ngu,
   nguon_du_lieu, ly_thuyet (union therapy), nam_xuat_hien, n_bai_gioi_thieu/
   danh_gia/phat_trien/su_dung, bai_gioi_thieu, link, availability.
4. Danh sach rieng MXH-thong-ke (nam + ten + bai) de Vy dem theo nam.
5. Xuat wave3_dataset_analysis.xlsx: sheet muc_bai / danh_muc_dataset / MXH_thong_ke.
Chay: python wave3_ds_merge.py <scratch_dir>"""
import csv
import json
import os
import re
import sys
import unicodedata
from collections import Counter, defaultdict

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
NOTES = os.path.dirname(os.path.abspath(__file__))
EXPORTS = os.path.join(os.path.dirname(NOTES), "bai_bao", "search_exports")
SRC = os.path.join(EXPORTS, "wave3_ai_final.csv")
SCRATCH = sys.argv[1] if len(sys.argv) > 1 else os.path.join(NOTES, "_ai_screen")
DS = os.path.join(SCRATCH, "ds_extract")
VAI = {"gioi-thieu-moi", "danh-gia", "phat-trien", "su-dung", "khong"}


def norm_ten(t):
    s = unicodedata.normalize("NFKD", t or "")
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    return re.sub(r"[^a-z0-9]+", "", s)


def main():
    import glob
    with open(os.path.join(DS, "ds_manifest.json"), encoding="utf-8") as f:
        manifest = json.load(f)
    ket, thieu_lo = {}, []
    for name, ids in sorted(manifest.items()):
        out = os.path.join(DS, name.replace(".jsonl", "_out.jsonl"))
        if not os.path.exists(out):
            thieu_lo.append(name)
            continue
        got = {}
        with open(out, encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                try:
                    d = json.loads(line)
                except json.JSONDecodeError:
                    continue
                if d.get("rec_id") in set(ids):
                    got[d["rec_id"]] = d
        if len(got) < len(ids):
            thieu_lo.append(f"{name}(thieu {len(ids)-len(got)})")
        ket.update(got)
    print(f"Phan quyet doc duoc: {len(ket)}/{sum(len(v) for v in manifest.values())} "
          f"| lo thieu/do dang: {len(thieu_lo)}")
    if thieu_lo:
        print("  ->", ", ".join(thieu_lo[:12]))

    with open(SRC, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    assert len(rows) == 7578

    n_cap_nhat = 0
    for r in rows:
        for c in ("lien_quan_dataset", "nhom_mien", "phuong_thuc", "don_da_luot",
                  "ngon_ngu", "nguon_du_lieu", "ds_minh_chung",
                  "dataset_cu_heuristic", "therapy_cu_heuristic"):
            r.setdefault(c, "")
        v = ket.get(r["rec_id"])
        if not v:
            continue
        n_cap_nhat += 1
        vai = (v.get("lien_quan_dataset") or "khong").strip()
        r["lien_quan_dataset"] = vai if vai in VAI else "khong"
        r["nhom_mien"] = (v.get("nhom_mien") or "").strip()
        r["phuong_thuc"] = (v.get("phuong_thuc") or "").strip()
        r["don_da_luot"] = (v.get("don_da_luot") or "").strip()
        r["ngon_ngu"] = (v.get("ngon_ngu") or "").strip()
        r["nguon_du_lieu"] = (v.get("nguon_du_lieu") or "").strip()
        r["ds_minh_chung"] = (v.get("minh_chung") or "").strip()
        ten_moi = (v.get("ds_ten") or "").strip()
        if ten_moi and ten_moi != r.get("dataset", ""):
            if not r["dataset_cu_heuristic"]:
                r["dataset_cu_heuristic"] = r.get("dataset", "")
            r["dataset"] = ten_moi
        th_moi = (v.get("therapy_ra_soat") or "").strip()
        if th_moi and th_moi != "(khong)" and th_moi != r.get("therapy", ""):
            if not r["therapy_cu_heuristic"]:
                r["therapy_cu_heuristic"] = r.get("therapy", "")
            r["therapy"] = th_moi.replace(",", ";")
        elif th_moi == "(khong)" and r.get("therapy", ""):
            if not r["therapy_cu_heuristic"]:
                r["therapy_cu_heuristic"] = r.get("therapy", "")
            r["therapy"] = ""

    with open(SRC, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    print(f"Da cap nhat {n_cap_nhat} dong vao {os.path.basename(SRC)}")

    # ------------------ tong hop muc DATASET ------------------
    ds_map = defaultdict(lambda: dict(ten_dem=Counter(), bai=[], vai=Counter(),
                                      mien=Counter(), pt=Counter(), luot=Counter(),
                                      ngon=Counter(), nguon=Counter(), ly=Counter(),
                                      nam=[], link=set(), avail=Counter()))
    for r in rows:
        vai = r.get("lien_quan_dataset") or ""
        if vai in ("", "khong"):
            continue
        for ten in [t.strip() for t in (r.get("dataset") or "").split(";") if t.strip()]:
            k = norm_ten(ten)
            if len(k) < 3:
                continue
            d = ds_map[k]
            d["ten_dem"][ten] += 1
            d["bai"].append((r["rec_id"], vai, r.get("year", ""), r.get("title", "")[:90]))
            d["vai"][vai] += 1
            for f_, key in [("nhom_mien", "mien"), ("phuong_thuc", "pt"),
                            ("don_da_luot", "luot"), ("ngon_ngu", "ngon"),
                            ("nguon_du_lieu", "nguon"), ("availability", "avail")]:
                val = (r.get(f_) or "").strip()
                if val and val not in ("khong-ro", "unstated"):
                    d[key][val] += 1
            for t2 in (r.get("therapy") or "").split(";"):
                if t2.strip():
                    d["ly"][t2.strip()] += 1
            if r.get("year"):
                d["nam"].append(r["year"])
            for u in (r.get("link_data") or "").split("; "):
                if u:
                    d["link"].add(u)

    def mode(c):
        return c.most_common(1)[0][0] if c else ""

    danh_muc = []
    for k, d in ds_map.items():
        gioi_thieu = [b for b in d["bai"] if b[1] in ("gioi-thieu-moi", "phat-trien")]
        # ten dai dien: bien the pho bien nhat; hoa het chu thuong -> uu tien co chu hoa
        ten_dd = sorted(d["ten_dem"].items(),
                        key=lambda kv: (-kv[1], kv[0] == kv[0].lower(), len(kv[0])))[0][0]
        danh_muc.append(dict(
            ten_dataset=ten_dd, nhom_mien=mode(d["mien"]), phuong_thuc=mode(d["pt"]),
            don_da_luot=mode(d["luot"]), ngon_ngu=mode(d["ngon"]),
            nguon_du_lieu=mode(d["nguon"]),
            ly_thuyet="; ".join(sorted(d["ly"])) or "",
            nam_som_nhat=min(d["nam"]) if d["nam"] else "",
            n_bai=len(d["bai"]), n_gioi_thieu=d["vai"]["gioi-thieu-moi"],
            n_danh_gia=d["vai"]["danh-gia"], n_phat_trien=d["vai"]["phat-trien"],
            n_su_dung=d["vai"]["su-dung"],
            bai_gioi_thieu="; ".join(f"{b[0]}({b[2]})" for b in gioi_thieu[:4]),
            link="; ".join(sorted(d["link"])[:3]), availability=mode(d["avail"]),
            cac_bai="; ".join(b[0] for b in d["bai"][:20])))
    danh_muc.sort(key=lambda x: (x["nhom_mien"], -(x["n_bai"])))

    mxh = [dict(rec_id=r["rec_id"], year=r.get("year", ""),
                ds_ten=r.get("dataset", ""), title=r["title"],
                minh_chung=r.get("ds_minh_chung", ""))
           for r in rows if (r.get("nhom_mien") or "") == "MXH-thong-ke"]
    mxh.sort(key=lambda x: x["year"])

    # ------------------ xuat XLSX ------------------
    from openpyxl import Workbook
    from openpyxl.utils.exceptions import IllegalCharacterError  # noqa: F401
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    def sach(v):
        return ILLEGAL_CHARACTERS_RE.sub("", str(v))[:32000]

    # cot heuristic da bi vong ra soat thay the -> bo khoi XLSX (CSV van giu du)
    BO_COT = {"priority_score", "band", "R_goi_y", "hint", "bucket",
              "dataset_da_biet", "dataset_ung_vien", "minhchung_therapy",
              "dataset_cu_heuristic", "therapy_cu_heuristic"}
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "muc_bai"
    cols1 = [c for c in rows[0].keys() if c not in BO_COT]
    ws1.append(cols1)
    for r in rows:
        ws1.append([sach(r.get(c) or "") for c in cols1])
    ws2 = wb.create_sheet("danh_muc_dataset")
    if danh_muc:
        cols2 = list(danh_muc[0].keys())
        ws2.append(cols2)
        for d in danh_muc:
            ws2.append([sach(d.get(c, "")) for c in cols2])
    ws3 = wb.create_sheet("MXH_thong_ke")
    if mxh:
        cols3 = list(mxh[0].keys())
        ws3.append(cols3)
        for d in mxh:
            ws3.append([sach(d.get(c, "")) for c in cols3])
    out_x = os.path.join(EXPORTS, "wave3_dataset_analysis.xlsx")
    wb.save(out_x)

    c_mien = Counter(x["nhom_mien"] for x in danh_muc)
    print(f"\n=== TONG HOP MUC DATASET ===")
    print(f"  So dataset duy nhat: {len(danh_muc)}")
    print(f"  Theo nhom mien: {dict(c_mien.most_common())}")
    print(f"  Danh sach MXH (thong ke rieng): {len(mxh)} bai")
    print(f"  XLSX: {out_x}")


if __name__ == "__main__":
    main()
